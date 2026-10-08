"""Slide and notebook pipeline for the Faculty Twin indexer (Block 1).

Reads the private lecture archive and writes private build outputs. Nothing this
script reads or writes belongs in the repo.

Inputs, per session folder ``~/Lecture Archive/<course>/2026 Fall/<NN> <date> <title>/``:
  slides.pdf    the most current deck (see the folder's SOURCES.md)
  slides.pptx   the matching PowerPoint file, used only for speaker notes and titles
  notebooks/    zero or more .ipynb files

Outputs, under ``~/Lecture Archive/_build/``:
  slides/<course>/s<NN>/<slide_id>.webp        one per PDF page, 1600 px wide
  slides/<course>/s<NN>/<slide_id>-thumb.webp  320 px wide
  slides/<course>/s<NN>/slides.json            [{slide_id, course, session, slide_number,
                                                 title, text, notes, flags, notes_match, ...}]
  slides/<course>/s<NN>/deck.json              deck-level facts (sources, counts, match quality)
  slides/missing.json                          sessions with no renderable deck, with the reason
  code/<course>/s<NN>.json                     [{cell_id, source, markdown_above, ...}], outputs stripped

Privacy: rosters in ``_private/rosters/`` are read only to test whether a page
mentions a student. Names are never printed, logged or written; a page gets a
flag, and strong matches in the stored text are replaced with ``[student]``.

Per-slide ``flags`` (flag, never drop; downstream stages decide):
  student_names_possible        full roster name, Andrew ID or email in text, notes or image (OCR)
  student_surname_possible      a roster surname alone (often a public figure or cited author)
  student_first_name_possible   a roster first name alone (weak)
  in_the_news                   the student "AI in the News" presentation slides / assignment
  student_presentation_possible "presented by", "team members", "group N", ...
  copyright_notice              (c), "copyright", "all rights reserved", "reprinted", ...
  third_party_source            "Source:", "et al", "adapted from", publisher or analyst names
  little_text                   image-only page; scrubbed OCR text is kept in ``ocr_text``
  notes_unmatched               the pptx exists but this page could not be matched to a slide
  no_clips_private_case         session whose class video must not be cut into clips
  secret_redacted               an API-key-like string was replaced with [REDACTED_KEY]
  pg_language                   cursing in the text, title, notes or OCR text was swapped for a mild
                                word (indexer/pg_filter.py); the slide image itself is unchanged

``notes_match.confidence``: high (text similarity >= 0.5), medium (>= 0.2),
positional (no text to compare, but the page sits in an unbroken one-to-one run
between text anchors), low / unmatched (no notes attached), no_pptx.

OCR uses Apple Vision through ``indexer/ocr_vision.swift`` (compiled once into
``_build/.cache``); without macOS it is skipped and flags use PDF text only.

Run (no .venv in the repo):
  uv run --no-project --with python-pptx --with pillow --with pypdf --with nbformat \\
      python indexer/slides.py [--course 70445] [--session 6] [--force]

The script is idempotent: pages whose images exist are not re-rendered, and a
deck whose sources and pipeline version are unchanged is not re-extracted
unless ``--force`` is given.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path

if __package__ in (None, ""):  # run as a script (python indexer/slides.py): make `indexer` importable
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from indexer import layout  # noqa: E402
from indexer.layout import NO_CLIP_SESSIONS, TERM, file_fingerprint, write_json  # noqa: E402
from indexer.pg_filter import smooth as pg_smooth  # noqa: E402
from indexer.roster import lowercase_dictionary_words, names_pattern, read_people  # noqa: E402
from indexer.roster import person as roster_person  # noqa: E402

# 8: PG filter on text, title, notes and OCR text (re-extracts text, never re-renders)
# 9: rosters with other column names (name/login_id, Student, SIS Login ID) join the name scrub
PIPELINE_VERSION = 9

# A module attribute (not read through layout at call time) so tests can point it at a fixture archive.
ARCHIVE = layout.DEFAULT_ARCHIVE
COURSES = layout.COURSE_FOLDERS

IMAGE_WIDTH = 1600
THUMB_WIDTH = 320
WEBP_QUALITY = 80
THUMB_QUALITY = 75

STUDENT = "[student]"


# --------------------------------------------------------------------------- paths


def build_dir() -> Path:
    """The private build folder under the current ARCHIVE."""
    return layout.build_dir(ARCHIVE)


def slide_id(course: str, session: int, n: int) -> str:
    """The stable id of page `n` of a session's deck, e.g. "70445-s06-012"."""
    return f"{course}-s{session:02d}-{n:03d}"


@dataclass
class Session:
    """One session folder of the archive."""

    course: str
    number: int
    date: str
    title: str
    folder: Path

    @property
    def tag(self) -> str:
        return f"{self.course} s{self.number:02d}"


def discover_sessions(course_filter: str | None, session_filter: int | None) -> list[Session]:
    """Session folders `<NN> <date> <title>` of the known courses, optionally filtered."""
    out = []
    for code, name in COURSES.items():
        if course_filter and code != course_filter.replace("-", ""):
            continue
        root = ARCHIVE / name / TERM
        for folder in sorted(root.glob("[0-9][0-9] *")):
            m = re.match(r"(\d{2}) (\d{4}-\d{2}-\d{2}) (.+)$", folder.name)
            if not m or not folder.is_dir():
                continue
            num = int(m.group(1))
            if session_filter and num != session_filter:
                continue
            out.append(Session(code, num, m.group(2), m.group(3), folder))
    return out


def inventory_notes() -> dict[tuple[str, int], str]:
    """The inventory's note per session (why a deck is missing), for missing.json."""
    data = layout.read_json(ARCHIVE / "_inventory" / "f26_inventory.json", {}) or {}
    notes = {}
    for name, course in data.get("courses", {}).items():
        code = name.split()[0].replace("-", "")
        for s in course.get("sessions", []):
            notes[(code, int(s["number"]))] = s.get("notes", "")
    return notes


# --------------------------------------------------------------------------- roster scrub


class NameScrubber:
    """Tests text for roster names without ever exposing them.

    Strong matches (full name, Andrew ID, email) are replaced with ``[student]``.
    A roster surname or first name on its own (ignoring ordinary English words)
    is only flagged: alone they collide with public figures and cited authors.
    """

    def __init__(self, rows: list[dict], dictionary: set[str] | None = None):
        dictionary = dictionary or set()
        self.full: set[str] = set()
        self.ids: set[str] = set()
        self.surnames: set[str] = set()
        self.firsts: set[str] = set()
        for r in rows:
            p = roster_person(r)  # any roster header layout (indexer/roster.py)
            if p is None:
                continue
            last, first, aid, email = p["last"], p["first"], p["andrew_id"], p["email"]
            if first and last:
                self.full.add(_norm_name(f"{first} {last}"))
                self.full.add(_norm_name(f"{first.split()[0]} {last}"))
            for tok in re.split(r"[\s\-]+", last.lower()):
                if len(tok) >= 3 and tok not in dictionary:
                    self.surnames.add(tok)
            for tok in first.lower().split():
                if len(tok) >= 3 and tok not in dictionary:
                    self.firsts.add(tok)
            if aid:
                self.ids.add(aid)
            if email:
                self.ids.add(email)
                self.ids.add(email.split("@")[0])
        self.full.discard("")
        self._full_re = names_pattern(self.full)

    @classmethod
    def from_dir(cls, roster_dir: Path) -> NameScrubber:
        """A scrubber for every roster CSV in `roster_dir`, ignoring ordinary English words."""
        return cls(read_people(roster_dir), lowercase_dictionary_words())

    def check(self, text: str) -> tuple[str, list[str]]:
        """Return (scrubbed text, flags). Never returns or logs a matched name."""
        flags: set[str] = set()
        if not text:
            return text, []
        if self._full_re and self._full_re.search(text):
            text = self._full_re.sub(STUDENT, text)
            flags.add("student_names_possible")

        def sub_token(m: re.Match) -> str:
            tok = m.group(0)
            low = tok.lower().rstrip(".")
            if low in self.ids:
                flags.add("student_names_possible")
                return STUDENT + tok[len(low):]
            if low in self.surnames and tok[0].isupper():
                # Flag only: surnames alone collide with public figures (and cited authors).
                flags.add("student_surname_possible")
            if low in self.firsts and tok[0].isupper():
                flags.add("student_first_name_possible")
            return tok

        text = re.sub(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+|[A-Za-z][A-Za-z'\-]*[A-Za-z0-9]*", sub_token, text)
        return text, sorted(flags)


def _norm_name(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().lower()


# --------------------------------------------------------------------------- content flags

# API-key-shaped strings in slides and notebooks. canvas_import.py has its own, wider pattern
# with a different marker; both markers are in built outputs, so they are kept apart.
SECRET_RE = re.compile(
    r"(sk-[A-Za-z0-9_\-]{16,}|sk-ant-[A-Za-z0-9_\-]{16,}|AKIA[0-9A-Z]{16}|gh[po]_[A-Za-z0-9]{20,}"
    r"|eyJ[A-Za-z0-9_\-]{10,}\.[A-Za-z0-9_\-]{10,}\.[A-Za-z0-9_\-]{5,}|pa-[A-Za-z0-9_\-]{20,}"
    r"|AIza[0-9A-Za-z_\-]{30,}|hf_[A-Za-z0-9]{20,}|sk_[a-f0-9]{40,})"
)

CONTENT_RULES = [
    ("in_the_news", re.compile(r"\bA[Il1](?:\s+Methods)?\s+in\s+the\s+News\b", re.I)),
    (
        "student_presentation_possible",
        re.compile(
            r"\bpresented\s+by\b|\bteam\s+members?\b|\bgroup\s+\d+\b|\bour\s+team\b|\bstudent\s+presentations?\b",
            re.I,
        ),
    ),
    (
        "copyright_notice",
        re.compile(
            r"©|\(c\)\s*\d{4}|\bcopyright\b|all\s+rights\s+reserved|used\s+with\s+permission|\breprinted\b", re.I
        ),
    ),
    (
        "third_party_source",
        re.compile(
            r"\bsource\s*:|\bsources\s*:|\bimage\s+(?:credit|source)|\bcredit\s*:|\bcourtesy\s+of\b|\badapted\s+from\b"
            r"|\bet\s+al\b|\bgetty\b|\bshutterstock\b|\bharvard\s+business\b|\bhbr\.org\b|\bmckinsey\b|\bgartner\b"
            r"|\bpearson\b|\bo'reilly\b|\bwiley\b|\bcengage\b|\bmodern\s+approach\b|\barxiv\b|\bdoi\b",
            re.I,
        ),
    ),
]


def content_flags(text: str) -> list[str]:
    """The CONTENT_RULES flags whose pattern appears in `text`."""
    return [name for name, rx in CONTENT_RULES if rx.search(text or "")]


def redact_secrets(text: str) -> tuple[str, bool]:
    """(text with key-shaped strings replaced by [REDACTED_KEY], whether anything changed)."""
    new = SECRET_RE.sub("[REDACTED_KEY]", text or "")
    return new, new != (text or "")


# --------------------------------------------------------------------------- PDF


def pdf_page_count(pdf: Path) -> int:
    """Pages in a PDF: pypdf when it can read the file, else poppler's pdfinfo."""
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(pdf)).pages)
    except Exception:  # pypdf missing, or a PDF it cannot parse: pdfinfo is the fallback, not a skip
        out = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, check=True).stdout
        return int(re.search(r"^Pages:\s+(\d+)", out, re.M).group(1))


def pdf_page_texts(pdf: Path, n_pages: int) -> list[str]:
    """Text per page via ``pdftotext -layout``; falls back to one call per page."""
    out = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", str(pdf), "-"], capture_output=True, check=True)
    pages = out.stdout.decode("utf-8", "replace").split("\f")
    if len(pages) >= n_pages and all(p == "" for p in pages[n_pages:]):
        return [clean_text(p) for p in pages[:n_pages]]
    texts = []
    for i in range(1, n_pages + 1):
        r = subprocess.run(
            ["pdftotext", "-layout", "-enc", "UTF-8", "-f", str(i), "-l", str(i), str(pdf), "-"], capture_output=True
        )
        texts.append(clean_text(r.stdout.decode("utf-8", "replace")))
    return texts


def clean_text(s: str) -> str:
    """Page text with runs of spaces shortened and blank lines collapsed to one."""
    lines = [re.sub(r"[ \t ]{2,}", "  ", ln).strip() for ln in s.replace("\f", "").splitlines()]
    out, blank = [], False
    for ln in lines:
        if not ln:
            blank = bool(out)
            continue
        if blank:
            out.append("")
            blank = False
        out.append(ln)
    return "\n".join(out).strip()


def render_page(pdf: Path, page: int, full: Path, thumb: Path) -> bool:
    """Render one page to webp (+ thumbnail). Returns True if work was done."""
    if full.exists() and thumb.exists():
        return False
    from PIL import Image

    with tempfile.TemporaryDirectory() as td:
        prefix = Path(td) / "p"
        subprocess.run(
            ["pdftoppm", "-png", "-singlefile", "-f", str(page), "-l", str(page),
             "-scale-to-x", str(IMAGE_WIDTH), "-scale-to-y", "-1", str(pdf), str(prefix)],
            check=True, capture_output=True,
        )
        with Image.open(str(prefix) + ".png") as im:
            im = im.convert("RGB")
            if im.width != IMAGE_WIDTH:
                im = im.resize((IMAGE_WIDTH, round(im.height * IMAGE_WIDTH / im.width)), Image.LANCZOS)
            _save_webp(im, full, WEBP_QUALITY)
            th = im.resize((THUMB_WIDTH, round(im.height * THUMB_WIDTH / im.width)), Image.LANCZOS)
            _save_webp(th, thumb, THUMB_QUALITY)
    return True


def _save_webp(image, path: Path, quality: int) -> None:
    """Save through a temporary file, so an interrupted run never leaves a broken image that looks done."""
    tmp = path.with_suffix(".tmp.webp")
    image.save(tmp, "WEBP", quality=quality, method=4)
    os.replace(tmp, path)


LIBREOFFICE_APP = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")

# Office apps driven by AppleScript when LibreOffice is missing or fails: (name, app, script).
# `{src}` and `{dst}` are filled with quoted POSIX paths; `{{` is a literal brace.
APPLESCRIPT_CONVERTERS = [
    ("Microsoft PowerPoint", Path("/Applications/Microsoft PowerPoint.app"),
     'tell application "Microsoft PowerPoint"\n'
     '  set p to open (POSIX file "{src}")\n'
     '  delay 5\n'
     '  save active presentation in (POSIX file "{dst}") as save as PDF\n'
     '  close active presentation saving no\n'
     'end tell'),
    ("Keynote", Path("/Applications/Keynote.app"),
     'with timeout of 600 seconds\n'
     'tell application "Keynote"\n'
     '  set d to open (POSIX file "{src}")\n'
     '  export d to (POSIX file "{dst}") as PDF with properties {{PDF image quality:Best, skipped slides:false}}\n'
     '  close d saving no\n'
     'end tell\n'
     'end timeout'),
]


def convert_pptx_to_pdf(pptx: Path, out_pdf: Path) -> tuple[bool, str]:
    """Try LibreOffice, then PowerPoint, then Keynote (AppleScript). Returns (ok, how)."""
    if out_pdf.exists() and out_pdf.stat().st_size > 0:
        return True, "cached"
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice and LIBREOFFICE_APP.exists():
        soffice = str(LIBREOFFICE_APP)
    if soffice and _convert_with_libreoffice(soffice, pptx, out_pdf):
        return True, "libreoffice"
    tried = ["LibreOffice: " + ("failed" if soffice else "not installed")]
    for name, app, script in APPLESCRIPT_CONVERTERS:
        if not app.exists():
            tried.append(f"{name}: not installed")
            continue
        problem = _convert_with_applescript(script, pptx, out_pdf)
        if problem is None:
            return True, name.lower().replace("microsoft ", "")
        tried.append(f"{name}: {problem}")
    return False, "no converter succeeded (" + "; ".join(tried) + ")"


def _convert_with_libreoffice(soffice: str, pptx: Path, out_pdf: Path) -> bool:
    """Headless LibreOffice export. True when it wrote out_pdf."""
    with tempfile.TemporaryDirectory() as td:
        r = subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", td, str(pptx)],
                           capture_output=True, timeout=1800)
        produced = list(Path(td).glob("*.pdf"))
        if r.returncode == 0 and produced:
            shutil.move(str(produced[0]), out_pdf)
            return True
    return False


def _convert_with_applescript(script: str, pptx: Path, out_pdf: Path) -> str | None:
    """Run one AppleScript export. None when it wrote out_pdf, else why it did not."""
    tmp = out_pdf.with_suffix(".tmp.pdf")
    src = str(pptx).replace('"', '\\"')
    dst = str(tmp).replace('"', '\\"')
    try:
        r = subprocess.run(["osascript", "-e", script.format(src=src, dst=dst)],
                           capture_output=True, text=True, timeout=660)
    except subprocess.TimeoutExpired:
        return "AppleScript timed out (likely an unanswered macOS Automation prompt)"
    if r.returncode == 0 and tmp.exists() and tmp.stat().st_size > 0:
        os.replace(tmp, out_pdf)
        return None
    return f"AppleScript failed (exit {r.returncode})"


# --------------------------------------------------------------------------- PPTX


@dataclass
class PptxSlide:
    """One slide of the PowerPoint file, used for titles and speaker notes."""

    index: int  # 1-based position in the pptx
    hidden: bool
    title: str
    text: str
    notes: str


def _shape_texts(shape) -> list[str]:
    """Every text frame and table row in a shape, walking into groups."""
    out = []
    try:
        is_group = shape.shape_type == 6  # MSO_SHAPE_TYPE.GROUP
    except Exception:  # python-pptx raises on shape types it does not model; treat them as plain shapes
        is_group = False
    if is_group:
        for s in shape.shapes:
            out.extend(_shape_texts(s))
        return out
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        out.append(shape.text_frame.text.strip())
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append("  ".join(cells))
    return out


def read_pptx(pptx: Path) -> list[PptxSlide]:
    """Title, text, speaker notes and hidden state of every slide in a PowerPoint file."""
    from pptx import Presentation

    prs = Presentation(str(pptx))
    slides = []
    for i, s in enumerate(prs.slides, start=1):
        title = ""
        try:
            if s.shapes.title is not None and s.shapes.title.has_text_frame:
                title = s.shapes.title.text_frame.text.strip()
        except Exception:  # a malformed title placeholder: the PDF text supplies the title instead
            title = ""
        texts = []
        for sh in s.shapes:
            texts.extend(_shape_texts(sh))
        notes = ""
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            notes = s.notes_slide.notes_text_frame.text.strip()
        slides.append(PptxSlide(i, s._element.get("show") == "0", title, "\n".join(texts), notes))
    return slides


# --------------------------------------------------------------------------- alignment

_STOP = {
    "the", "a", "an", "and", "or", "of", "to", "in", "on", "for", "is", "are", "be", "with", "as", "by", "at",
    "it", "this", "that", "from",
}


def tokens(text: str) -> set[str]:
    """Lowercase content words of `text`, for comparing a PDF page with a pptx slide."""
    return {t for t in re.findall(r"[a-z0-9]+", (text or "").lower()) if len(t) > 1 and t not in _STOP}


def similarity(a: str, b: str) -> float | None:
    """Jaccard overlap of the two texts' words; None when neither has any words."""
    ta, tb = tokens(a), tokens(b)
    if not ta and not tb:
        return None
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)


MATCH_BONUS = 0.15
HIDDEN_PENALTY = 0.3


def align(pdf_texts: list[str], slides: list[PptxSlide]) -> list[tuple[int | None, float | None]]:
    """Monotonic alignment of PDF pages to pptx slides.

    Maximises the sum of (text similarity + a small bonus per matched pair -
    a penalty for hidden slides), with free gaps on both sides. Hidden slides
    are usually missing from the PDF export, so the penalty steers around them
    unless their text clearly matches a page. Returns, per PDF page, the
    0-based pptx index and its similarity (None when unmatched / no text).
    """
    n, m = len(pdf_texts), len(slides)
    sims = [[similarity(pdf_texts[i], slides[j].text) for j in range(m)] for i in range(n)]

    def w(i, j):
        s = sims[i][j]
        return (s if s is not None else 0.0) + MATCH_BONUS - (HIDDEN_PENALTY if slides[j].hidden else 0.0)

    dp = [[0.0] * (m + 1) for _ in range(n + 1)]
    for i in range(n - 1, -1, -1):
        for j in range(m - 1, -1, -1):
            dp[i][j] = max(dp[i + 1][j], dp[i][j + 1], w(i, j) + dp[i + 1][j + 1])
    out: list[tuple[int | None, float | None]] = [(None, None)] * n
    i = j = 0
    while i < n and j < m:
        if dp[i][j] == w(i, j) + dp[i + 1][j + 1] and w(i, j) > 0:
            out[i] = (j, sims[i][j])
            i += 1
            j += 1
        elif dp[i][j] == dp[i + 1][j]:
            i += 1
        else:
            j += 1
    return out


HIGH_MATCH = 0.5  # text similarity for a "high" confidence notes match
MEDIUM_MATCH = 0.2  # ... for "medium"; also what makes a page an anchor for positional matching


def confidence(sim: float | None, positional_ok: bool) -> str:
    """The notes_match.confidence label for one page (see the module docstring)."""
    if sim is not None and sim >= HIGH_MATCH:
        return "high"
    if sim is not None and sim >= MEDIUM_MATCH:
        return "medium"
    return "positional" if positional_ok else "low"


def positional_ok(alignment: list[tuple[int | None, float | None]], slides: list[PptxSlide]) -> list[bool]:
    """True for a weakly matched page that sits in an unbroken run between text anchors.

    Image-only pages have no text to compare. Such a page is still trusted when
    the pages between its two nearest anchors (similarity >= 0.2, or the deck
    ends) line up one-to-one with the visible pptx slides between them.
    """
    n = len(alignment)
    rank, r = [], 0  # visible rank of each pptx slide
    for s in slides:
        rank.append(r)
        r += 0 if s.hidden else 1
    total_visible = r
    anchors = [(-1, -1)]
    for i, (j, sim) in enumerate(alignment):
        if j is not None and sim is not None and sim >= MEDIUM_MATCH and not slides[j].hidden:
            anchors.append((i, rank[j]))
    anchors.append((n, total_visible))
    ok = [False] * n
    for (pi, pr), (ni, nr) in zip(anchors, anchors[1:]):
        if ni - pi != nr - pr:
            continue
        for i in range(pi + 1, ni):
            j = alignment[i][0]
            ok[i] = j is not None and not slides[j].hidden and rank[j] - pr == i - pi
    return ok


# --------------------------------------------------------------------------- titles


def boilerplate_lines(page_texts: list[str]) -> set[str]:
    """Lines repeated on at least 30% of pages (footers, course names, page furniture)."""
    counts: Counter = Counter()
    for t in page_texts:
        counts.update({ln.strip().lower() for ln in t.splitlines() if ln.strip()})
    limit = max(3, int(0.3 * len(page_texts)))
    return {ln for ln, c in counts.items() if c >= limit}


def strip_boilerplate(text: str, boiler: set[str]) -> str:
    """`text` without the boilerplate lines."""
    return "\n".join(ln for ln in text.splitlines() if ln.strip().lower() not in boiler).strip()


_WRAP_RE = re.compile(r"(\b(?:and|or|of|for|the|to|in|with|a|an|on|vs\.?)|[,:&\-])$", re.I)
_URL_RE = re.compile(r"^(https?://|www\.)\S+$", re.I)


def pick_title(text: str, pptx_title: str = "") -> str:
    """The pptx title placeholder when there is one, else the first meaningful PDF line."""
    if pptx_title and re.search(r"[A-Za-z]", pptx_title) and not _URL_RE.match(pptx_title.strip()):
        return _short(pptx_title)
    lines = [ln.strip() for ln in text.splitlines()]
    for i, ln in enumerate(lines):
        if len(ln) < 3 or not re.search(r"[A-Za-z]{2}", ln) or _URL_RE.match(ln):
            continue
        if re.fullmatch(r"(page\s*)?\d+(\s*(/|of)\s*\d+)?", ln, re.I):
            continue
        title = re.split(r"\s{2,}", ln)[0]
        # A title wrapped onto a second line ("AI Methods for Social and" / "Visual Data").
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        if nxt and _WRAP_RE.search(title) and title == ln:
            title += " " + re.split(r"\s{2,}", nxt)[0]
        return _short(title)
    return ""


def _short(s: str, limit: int = 120) -> str:
    s = re.sub(r"\s+", " ", s).strip()
    return s if len(s) <= limit else s[: limit - 1].rsplit(" ", 1)[0] + "…"


# --------------------------------------------------------------------------- deck


@dataclass
class Deck:
    """One session's deck being processed: its sources, output folder, and what this run did."""

    session: Session
    pdf: Path
    pptx: Path | None
    source: str
    out: Path = None
    pages: int = 0
    rendered: int = 0
    skipped_extract: bool = False
    ocr_used: bool = False
    quality: dict = field(default_factory=dict)
    flags: dict = field(default_factory=dict)

    def __post_init__(self):
        self.out = build_dir() / "slides" / self.session.course / f"s{self.session.number:02d}"

    def sid(self, p: int) -> str:
        return slide_id(self.session.course, self.session.number, p)

    def image(self, p: int) -> Path:
        return self.out / f"{self.sid(p)}.webp"


def render_deck(deck: Deck, workers: int) -> None:
    """Render every page that is not on disk yet (resumable per page)."""
    deck.out.mkdir(parents=True, exist_ok=True)
    deck.pages = pdf_page_count(deck.pdf)
    jobs = [(deck.pdf, p, deck.image(p), deck.out / f"{deck.sid(p)}-thumb.webp") for p in range(1, deck.pages + 1)]
    with ThreadPoolExecutor(max_workers=workers) as ex:
        deck.rendered = sum(ex.map(lambda a: render_page(*a), jobs))


def roster_fingerprint(roster_dir: Path | None = None) -> str:
    """A hash of the roster files' names, sizes and times (never their contents' text).

    Part of the extraction cache key: a roster added or changed after a deck was
    extracted must re-run the name scrub on that deck. (Oct 5: rosters from other
    terms were added after the decks were extracted, and the cache kept a slide
    that names a student, because the key only covered the PDF and pptx.)
    """
    roster_dir = roster_dir or ARCHIVE / "_private" / "rosters"
    files = sorted(roster_dir.glob("*.csv")) if roster_dir.is_dir() else []
    return hashlib.sha256(json.dumps(file_fingerprint(*files)).encode()).hexdigest()[:16]


def extraction_current(deck: Deck, roster_fp: str | None = None) -> bool:
    """True when slides.json was built from the same sources and rosters by this pipeline version."""
    deck_path, slides_path = deck.out / "deck.json", deck.out / "slides.json"
    if not (deck_path.exists() and slides_path.exists()):
        return False
    try:
        old = json.loads(deck_path.read_text())
    except ValueError:
        return False
    roster_fp = roster_fp if roster_fp is not None else roster_fingerprint()
    unchanged = (
        old.get("pipeline_version") == PIPELINE_VERSION
        and old.get("fingerprint") == file_fingerprint(deck.pdf, deck.pptx)
        and old.get("pages") == deck.pages
        and old.get("roster_fingerprint") == roster_fp
    )
    if unchanged:
        deck.quality = old.get("notes_match", {})
        deck.flags = old.get("flag_counts", {})
        deck.ocr_used = bool(old.get("ocr"))
    return unchanged


@dataclass
class PageText:
    """One page's stored text fields, after every privacy filter, and the flags they raised."""

    text: str
    title: str
    notes: str
    ocr_text: str
    flags: set[str]


def clean_page_text(scrubber: NameScrubber, text: str, title: str, notes: str, ocr_raw: str,
                    pg_counts: Counter) -> PageText:
    """Run a page's text through the filters, in privacy order: names, then keys, then the PG rule.

    Content flags read the raw text first, so a flag never depends on what a filter removed.
    The title gets no key redaction: it is a short heading, and the stored outputs were built that way.
    """
    flags: set[str] = set(content_flags("\n".join((text, notes, ocr_raw))))
    text, f1 = scrubber.check(text)
    title, f2 = scrubber.check(title)
    notes, f3 = scrubber.check(notes)
    ocr_text, f4 = scrubber.check(ocr_raw)
    flags.update(f1, f2, f3, f4)
    text, s1 = redact_secrets(text)
    notes, s2 = redact_secrets(notes)
    ocr_text, s3 = redact_secrets(ocr_text)
    if s1 or s2 or s3:
        flags.add("secret_redacted")
    # PG rule: mild words for cursing in the stored text (the slide image is unchanged)
    text, p1 = pg_smooth(text, pg_counts)
    title, p2 = pg_smooth(title, pg_counts)
    notes, p3 = pg_smooth(notes, pg_counts)
    ocr_text, p4 = pg_smooth(ocr_text, pg_counts)
    if p1 or p2 or p3 or p4:
        flags.add("pg_language")
    return PageText(text, title, notes, ocr_text, flags)


def _slide_record(deck: Deck, p: int, page: PageText, pptx_slide: PptxSlide | None, sim: float | None,
                  conf: str) -> dict:
    """The slides.json row for page `p`."""
    sess, sid = deck.session, deck.sid(p)
    rec = {
        "slide_id": sid,
        "course": sess.course,
        "session": sess.number,
        "slide_number": p,
        "title": page.title,
        "text": page.text,
        "notes": page.notes,
        "flags": sorted(page.flags),
        "notes_match": {
            "pptx_slide": pptx_slide.index if pptx_slide else None,
            "similarity": None if sim is None else round(sim, 3),
            "confidence": conf,
        },
        "image": f"slides/{sess.course}/s{sess.number:02d}/{sid}.webp",
        "thumb": f"slides/{sess.course}/s{sess.number:02d}/{sid}-thumb.webp",
    }
    if "little_text" in page.flags and page.ocr_text.strip():
        # Image-only slide: keep the (scrubbed) OCR text so the page can still be found.
        rec["ocr_text"] = page.ocr_text.strip()
        if not rec["title"]:
            rec["title"] = pick_title(rec["ocr_text"])
    return rec


def extract_deck(deck: Deck, scrubber: NameScrubber, ocr: dict[str, str] | None) -> None:
    """Write slides.json and deck.json: text, title, notes, match quality and flags per page."""
    sess, n = deck.session, deck.pages
    raw = pdf_page_texts(deck.pdf, n)
    boiler = boilerplate_lines(raw)
    texts = [strip_boilerplate(t, boiler) for t in raw]

    pslides = read_pptx(deck.pptx) if deck.pptx else []
    visible = [s for s in pslides if not s.hidden]
    alignment = align(texts, pslides) if pslides else [(None, None)] * n
    positional = positional_ok(alignment, pslides) if pslides else [False] * n

    records, qual, flag_counts = [], {}, {}
    pg_counts: Counter = Counter()
    for p in range(1, n + 1):
        j, sim = alignment[p - 1]
        ps = pslides[j] if j is not None else None
        conf = confidence(sim, positional[p - 1]) if ps else ("unmatched" if pslides else "no_pptx")
        use_pptx = ps is not None and conf in ("high", "medium", "positional")
        notes = ps.notes if use_pptx else ""
        title = pick_title(texts[p - 1], ps.title if use_pptx else "")
        ocr_raw = (ocr or {}).get(str(deck.image(p)), "")

        page = clean_page_text(scrubber, texts[p - 1], title, notes, ocr_raw, pg_counts)
        if len(re.sub(r"\W", "", page.text)) < 15:
            page.flags.add("little_text")
        if (sess.course, sess.number) in NO_CLIP_SESSIONS:
            page.flags.add("no_clips_private_case")
        if deck.source.startswith("converted"):
            page.flags.add("converted_from_pptx")
        if pslides and conf in ("low", "unmatched"):
            page.flags.add("notes_unmatched")

        qual[conf] = qual.get(conf, 0) + 1
        for f in page.flags:
            flag_counts[f] = flag_counts.get(f, 0) + 1
        records.append(_slide_record(deck, p, page, ps, sim, conf))

    write_json(deck.out / "slides.json", records)
    write_json(deck.out / "deck.json", {
        "pipeline_version": PIPELINE_VERSION,
        "course": sess.course,
        "session": sess.number,
        "date": sess.date,
        "session_title": sess.title,
        "source": deck.source,
        "pdf": deck.pdf.name,
        "pptx": deck.pptx.name if deck.pptx else None,
        "fingerprint": file_fingerprint(deck.pdf, deck.pptx),
        "roster_fingerprint": roster_fingerprint(),
        "pages": n,
        "pptx_slides": len(pslides),
        "pptx_hidden": len(pslides) - len(visible),
        "count_match": len(visible) == n if pslides else None,
        "notes_match": qual,
        "flag_counts": flag_counts,
        "pg_changes": dict(pg_counts),  # labels such as "hell -> heck", counts only
        "ocr": ocr is not None,
        "boilerplate_lines_removed": len(boiler),
    })
    deck.quality, deck.flags, deck.ocr_used = qual, flag_counts, ocr is not None


# --------------------------------------------------------------------------- OCR

OCR_SOURCE = Path(__file__).with_name("ocr_vision.swift")


def ocr_images(paths: list[Path]) -> dict[str, str] | None:
    """OCR slide images with Apple Vision (macOS). Returns None when unavailable.

    The helper is compiled once into the private build folder. Text is returned
    to the caller only; it is never printed.
    """
    if not paths:
        return {}
    swiftc = shutil.which("swiftc")
    if sys.platform != "darwin" or not swiftc or not OCR_SOURCE.exists():
        return None
    binary = build_dir() / ".cache" / "ocr_vision"
    if not binary.exists() or binary.stat().st_mtime < OCR_SOURCE.stat().st_mtime:
        binary.parent.mkdir(parents=True, exist_ok=True)
        r = subprocess.run([swiftc, "-O", "-o", str(binary), str(OCR_SOURCE)], capture_output=True)
        if r.returncode != 0:
            return None
    out: dict[str, str] = {}
    chunk = 400
    for i in range(0, len(paths), chunk):
        r = subprocess.run([str(binary), *map(str, paths[i:i + chunk])], capture_output=True, timeout=3600)
        if r.returncode != 0:
            return None
        out.update(json.loads(r.stdout or b"{}"))
    return out


# --------------------------------------------------------------------------- notebooks


def process_notebooks(sess: Session, scrubber: NameScrubber) -> tuple[int, int, dict]:
    """Write code/<course>/s<NN>.json: every code cell, scrubbed, with the markdown above it.

    Returns (notebooks read, cells written, flag counts). Outputs are never stored.
    """
    import nbformat

    nbs = sorted((sess.folder / "notebooks").glob("*.ipynb"))
    if not nbs:
        return 0, 0, {}
    cells_out, flag_counts = [], {}
    for k, nb_path in enumerate(nbs, start=1):
        try:
            nb = nbformat.read(str(nb_path), as_version=4)
        except Exception as e:  # malformed notebook: record and move on
            print(f"  ! {sess.tag} notebook {k}: unreadable ({type(e).__name__})", file=sys.stderr)
            continue
        md_buf: list[str] = []
        for idx, cell in enumerate(nb.cells):
            src = cell.get("source", "")
            if isinstance(src, list):
                src = "".join(src)
            if cell.cell_type == "markdown":
                if src.strip():
                    md_buf.append(src.strip())
                continue
            if cell.cell_type != "code" or not src.strip():
                continue
            row = _code_cell(sess, k, nb_path.name, idx, src, "\n\n".join(md_buf), scrubber)
            for f in row["flags"]:
                flag_counts[f] = flag_counts.get(f, 0) + 1
            cells_out.append(row)
            md_buf = []
    write_json(build_dir() / "code" / sess.course / f"s{sess.number:02d}.json", cells_out)
    return len(nbs), len(cells_out), flag_counts


def _code_cell(sess: Session, k: int, notebook: str, idx: int, src: str, markdown: str,
               scrubber: NameScrubber) -> dict:
    """One code/<course>/s<NN>.json row: the cell and the markdown above it, names and keys removed."""
    flags: set[str] = set()
    src, f1 = scrubber.check(src)
    md, f2 = scrubber.check(markdown)
    src, r1 = redact_secrets(src)
    md, r2 = redact_secrets(md)
    flags.update(f1, f2)
    if r1 or r2:
        flags.add("secret_redacted")
    return {
        "cell_id": f"{sess.course}-s{sess.number:02d}-nb{k}-c{idx:03d}",
        "course": sess.course,
        "session": sess.number,
        "notebook": notebook,
        "cell_index": idx,
        "source": src,
        "markdown_above": md,
        "flags": sorted(flags),
    }


# --------------------------------------------------------------------------- main


def dir_size(path: Path) -> int:
    """Total bytes of the files under `path`."""
    return sum(p.stat().st_size for p in path.rglob("*") if p.is_file())


def _missing(sess: Session, reason: str) -> dict:
    """A missing.json row: a session with no renderable deck, and why."""
    return {"course": sess.course, "session": sess.number, "date": sess.date, "title": sess.title, "reason": reason}


def _find_deck(sess: Session, args: argparse.Namespace, missing: dict, inventory: dict) -> Deck | None:
    """The session's renderable deck (converting a pptx-only deck if allowed), else note it in `missing`."""
    key = (sess.course, sess.number)
    pdf, pptx = sess.folder / "slides.pdf", sess.folder / "slides.pptx"
    source = "slides.pdf"
    if not pdf.exists() and pptx.exists() and not args.no_convert:
        conv = build_dir() / "slides" / sess.course / f"s{sess.number:02d}" / "source_converted.pdf"
        print(f"{sess.tag}: no PDF, converting pptx ...", flush=True)
        ok, how = convert_pptx_to_pdf(pptx, conv)
        if ok:
            pdf, source = conv, f"converted from slides.pptx ({how})"
        else:
            missing[key] = _missing(sess, f"pptx only, no PDF export; {how}")
    if pdf.exists():
        missing.pop(key, None)
        return Deck(sess, pdf, pptx if pptx.exists() else None, source)
    if args.no_convert and pptx.exists() and key not in missing:
        missing[key] = _missing(sess, "pptx only, no PDF export; conversion skipped (--no-convert)")
    elif key not in missing:
        note = inventory.get(key, "")
        missing[key] = _missing(sess, "no slide deck in the archive" + (f" (inventory: {note})" if note else ""))
    return None


def _extract_changed(results: list[Deck], scrubber: NameScrubber, force: bool, use_ocr: bool) -> None:
    """Text, notes and flags, only for decks whose sources, rosters or pipeline changed."""
    roster_fp = roster_fingerprint()
    todo = [d for d in results if force or not extraction_current(d, roster_fp)]
    for d in results:
        d.skipped_extract = d not in todo
    ocr = None
    if todo and use_ocr:
        imgs = [d.image(p) for d in todo for p in range(1, d.pages + 1)]
        print(f"OCR of {len(imgs)} slide images ...", flush=True)
        ocr = ocr_images(imgs)
        if ocr is None:
            print("  OCR unavailable (needs macOS + swiftc); name flags use PDF text and notes only")
    for d in todo:
        print(f"{d.session.tag}: extracting text, notes, flags ...", flush=True)
        extract_deck(d, scrubber, ocr)


def _print_summary(results: list[Deck], sessions: list[Session], missing: dict, code_rows: list,
                   started: float) -> None:
    """Counts per deck, flag totals, missing decks and notebooks. Never slide text."""
    print("\n=== Slide pipeline summary ===")
    print(f"{'deck':<11}{'pages':>6}{'new img':>8}  notes match (pages)                       flags")
    for r in results:
        q = ", ".join(f"{k} {v}" for k, v in sorted(r.quality.items()))
        f = ", ".join(f"{k} {v}" for k, v in sorted(r.flags.items()))
        tag = r.session.tag + (" *" if r.skipped_extract else "")
        print(f"{tag:<11}{r.pages:>6}{r.rendered:>8}  {q:<42}{f}")
    print(f"decks: {len(results)}, pages: {sum(r.pages for r in results)}, "
          f"newly rendered: {sum(r.rendered for r in results)}  (* = text unchanged, reused); "
          f"OCR used on {sum(1 for r in results if r.ocr_used)} decks")
    total_flags: dict[str, int] = {}
    for r in results:
        for k, v in r.flags.items():
            total_flags[k] = total_flags.get(k, 0) + v
    print("flag totals:", ", ".join(f"{k} {v}" for k, v in sorted(total_flags.items())) or "none")
    sel = {(s.course, s.number) for s in sessions}
    print("missing decks:")
    for m in sorted(missing.values(), key=lambda m: (m["course"], m["session"])):
        if (m["course"], m["session"]) in sel:
            print(f"  {m['course']} s{m['session']:02d}: {m['reason'][:110]}")
    print("notebooks:")
    for sess, n_nb, n_cells, fl in code_rows:
        f = ", ".join(f"{k} {v}" for k, v in sorted(fl.items()))
        print(f"  {sess.tag}: {n_nb} notebooks, {n_cells} code cells {('[' + f + ']') if f else ''}")
    slides_root, code_root = build_dir() / "slides", build_dir() / "code"
    print(f"output size: slides {dir_size(slides_root) / 1e6:.1f} MB, "
          f"code {dir_size(code_root) / 1e6 if code_root.exists() else 0:.1f} MB; {time.time() - started:.0f} s")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--course", choices=list(COURSES) + ["70-445", "45-884"])
    ap.add_argument("--session", type=int)
    ap.add_argument("--force", action="store_true", help="re-extract text/notes even if unchanged")
    ap.add_argument("--no-convert", action="store_true", help="do not try to convert pptx-only decks")
    ap.add_argument("--no-ocr", action="store_true", help="skip Apple Vision OCR of slide images")
    ap.add_argument("--workers", type=int, default=max(2, (os.cpu_count() or 4) - 2))
    args = ap.parse_args(argv)

    started = time.time()
    sessions = discover_sessions(args.course, args.session)
    if not sessions:
        print("no sessions matched", file=sys.stderr)
        return 1
    scrubber = NameScrubber.from_dir(ARCHIVE / "_private" / "rosters")
    inventory = inventory_notes()

    missing_path = build_dir() / "slides" / "missing.json"
    try:
        missing = {(m["course"], m["session"]): m for m in json.loads(missing_path.read_text())}
    except (OSError, ValueError):
        missing = {}

    results: list[Deck] = []
    code_rows = []
    for sess in sessions:
        deck = _find_deck(sess, args, missing, inventory)
        if deck is not None:
            print(f"{sess.tag}: rendering {deck.source} ...", flush=True)
            render_deck(deck, args.workers)
            results.append(deck)
        n_nb, n_cells, nb_flags = process_notebooks(sess, scrubber)
        if n_nb:
            code_rows.append((sess, n_nb, n_cells, nb_flags))

    write_json(missing_path, sorted(missing.values(), key=lambda m: (m["course"], m["session"])))
    _extract_changed(results, scrubber, args.force, use_ocr=not args.no_ocr)
    _print_summary(results, sessions, missing, code_rows, started)
    return 0


if __name__ == "__main__":
    sys.exit(main())
